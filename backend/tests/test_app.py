import io
import time
import zipfile
import pytest
from fastapi.testclient import TestClient
from docx import Document
from reportlab.pdfgen import canvas
from backend import storage
from backend.analysis import analyze, segment, AbstainingDetector
from backend.main import app, jobs, requests

@pytest.fixture
def client(tmp_path, monkeypatch):
    monkeypatch.setattr(storage,'DB',tmp_path/'test.sqlite3')
    jobs.clear(); requests.clear()
    with TestClient(app) as c: yield c

def completed(c,text='The birds return. The birds sing in spring. I watch them from my window.',**kw):
    response=c.post('/api/jobs',json={'text':text,**kw})
    assert response.status_code==202,response.text
    id=response.json()['id']
    for _ in range(100):
        j=c.get('/api/jobs/'+id).json()
        if j['status'] in ('completed','failed'): return id,j
        time.sleep(.01)
    pytest.fail('Analysis did not complete')

def saved(c,**kw):
    id,j=completed(c,**kw)
    assert j['status']=='completed'
    assert c.post('/api/jobs/'+id+'/save').status_code==200
    return id

def test_features_and_segmentation():
    r=analyze('Cat cat dog. Bird flies!')
    assert r['word_count']==5
    assert r['features']['lexical_diversity']==.8
    assert r['features']['mean_sentence_words']==2.5
    assert r['features']['sentence_stddev']==.5
    assert len(segment('Dr. Smith measured 3.14 units. It worked.'))==2
    assert r['detector']['probability'] is None
    assert r['detector']['status']=='abstained'
    assert any('Short sample' in w for w in r['warnings'])

@pytest.mark.parametrize('text',[
    'I remember the bus ride home. My coat was soaked, and I laughed at the weather.',
    'Furthermore, sustainable development requires a comprehensive approach. Moreover, the framework enables effective collaboration.',
    'I wrote the opening. Furthermore, this framework addresses the issue. Then I rewrote the conclusion myself.',
    'Hi there.',
    'She said, “These words are quoted.” The essay includes a citation (Smith, 2024).',
    'I am learning English. Sometimes my sentences is not correct but I keep practicing.'
])
def test_scenarios_abstain_without_labels(text):
    r=analyze(text)
    assert r['word_count']>0 and r['sentences']
    assert r['detector']['probability'] is None

def test_detector_interface():
    class EvaluatedAdapter:
        def predict(self,text,language): return {'status':'abstained','probability':None,'reason':'Outside evaluated domain'}
    assert analyze('A sentence.',detector=EvaluatedAdapter())['detector']['reason']=='Outside evaluated domain'

def test_language_and_failed_job(client):
    id,j=completed(client,language='fr')
    assert j['status']=='failed' and 'English' in j['error']
    assert client.post('/api/jobs/'+id+'/save').status_code==409
    with pytest.raises(ValueError): analyze('这是一个中文句子。')

def test_uploads(client):
    assert client.post('/api/extract',files={'file':('x.txt',b'Hello\n\nWorld')}).json()['text']=='Hello\n\nWorld'
    doc=Document(); doc.add_paragraph('First paragraph.'); table=doc.add_table(rows=1,cols=2);table.cell(0,0).text='Table text'; doc.add_paragraph('Last paragraph.')
    buf=io.BytesIO();doc.save(buf)
    text=client.post('/api/extract',files={'file':('x.docx',buf.getvalue())}).json()['text']
    assert text.index('First')<text.index('Table')<text.index('Last')
    buf=io.BytesIO();pdf=canvas.Canvas(buf);pdf.drawString(20,700,'A real PDF document.');pdf.save()
    result=client.post('/api/extract',files={'file':('x.pdf',buf.getvalue())})
    assert result.status_code==200 and 'real PDF' in result.json()['text']
    assert client.post('/api/extract',files={'file':('x.exe',b'bad')}).status_code==415
    assert client.post('/api/extract',files={'file':('x.pdf',b'bad')}).status_code==422
    buf=io.BytesIO();p=canvas.Canvas(buf);p.showPage();p.save()
    assert 'OCR' in client.post('/api/extract',files={'file':('scan.pdf',buf.getvalue())}).json()['detail']
    assert client.post('/api/extract',files={'file':('x.txt',b'a'*(10*1024*1024+1))}).status_code==413

def test_persistence_profile_comparison_and_reports(client):
    a=saved(client,title='First draft'); b=saved(client,title='Next draft',text='The birds fly away. I close my window.',group_id=a,notes='New ending')
    assert len(client.get('/api/analyses').json())==2
    assert client.get('/api/analyses/'+a).json()['text'].startswith('The birds')
    assert client.put('/api/profile',json={'consent':False,'analysis_ids':[a]}).status_code==422
    assert len(client.put('/api/profile',json={'consent':True,'analysis_ids':[a]}).json()['references'])==1
    d=client.get('/api/compare',params={'left':a,'right':b}).json()
    assert d['diff'] and d['changes']['mean_sentence_words']!=0
    for fmt in ('pdf','json','txt'):
        r=client.get(f'/api/analyses/{a}/export/{fmt}'); assert r.status_code==200
        if fmt=='pdf': assert r.content.startswith(b'%PDF-')
        elif fmt=='json': assert r.json()['title']=='First draft'
        else: assert 'First draft' in r.text and 'unavailable' in r.text
    assert '3 sentences' in client.post(f'/api/analyses/{a}/assistant',json={'topic':'variation'}).json()['answer']
    client.delete('/api/analyses/'+a)
    assert client.get('/api/analyses/'+a).status_code==404
    assert client.get('/api/profile').json()['references']==[]
    client.delete('/api/profile');assert client.get('/api/profile').json()['consent'] is None
    client.delete('/api/data');assert client.get('/api/analyses').json()==[]

def test_security_retention_and_cancellation(client):
    assert client.post('/api/jobs',json={'text':'x'},headers={'Origin':'https://evil.example'}).status_code==403
    assert client.get('/api/health').headers['x-content-type-options']=='nosniff'
    assert client.post('/api/jobs',json={'text':'x'*200001}).status_code==422
    a=saved(client)
    with storage.connect() as db: db.execute("UPDATE analyses SET created='2000-01-01T00:00:00' WHERE id=?",(a,))
    assert client.put('/api/settings',json={'retention_days':30}).status_code==200
    assert client.get('/api/analyses').json()==[]
    id,j=completed(client)
    assert client.delete('/api/jobs/'+id).status_code==204
    assert client.get('/api/jobs/'+id).status_code==404
